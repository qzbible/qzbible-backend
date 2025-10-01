import io
from openpyxl import Workbook
from datetime import datetime
#__________________________________________________________
#                                                          |
# Exportations des données du manager sous format excel    |
#__________________________________________________________|
def export_produits_to_excel(produits):
    wb = Workbook()
    ws = wb.active
    ws.title = "Produits"

    # En-têtes
    headers = [
        "Nom", "Description", "Prix", "Unité", "Quantité",
        "État", "Catégorie", "Image", "Date de création"
    ]
    ws.append(headers)

    for p in produits:
        ws.append([
            p.get("nom"),
            p.get("description"),
            p.get("prix"),
            p.get("unite"),
            p.get("quantite"),
            p.get("etat"),
            p.get("categorie_nom"),
            p.get("image_url"),
            p.get("created_at").strftime("%Y-%m-%d %H:%M") if p.get("created_at") else ""
        ])

    # Sauvegarde dans un fichier en mémoire
    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    return file_stream




def export_clients_to_excel(clients):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clients"

    headers = [
        "Nom", "Téléphone", "Quartier", "Ville",
        "Nombre de Livraisons", "Montant Total",
        "Créé par", "Créé le"
    ]
    ws.append(headers)

    for c in clients:
        ws.append([
            c.get("nom", ""),
            c.get("telephone", ""),
            c.get("quartier", ""),
            c.get("ville", ""),
            c.get("nombre_livraisons", 0),
            c.get("montant_total", 0),
            f"{c.get('created_by_name', '')} {c.get('created_by_first_name', '')}",
            c.get("created_at").strftime("%Y-%m-%d %H:%M") if c.get("created_at") else ""
        ])

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    return file_stream




def export_livraisons_to_excel(livraisons):
    from openpyxl import Workbook
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = "Livraisons"

    headers = [
        "Client", "Téléphone", "Ville", "Quartier",
        "Montant Total", "Montant Payé", "Montant Restant",
        "Quantité", "Statut Paiement", "Statut Livraison",
        "Date Commande", "Créé le"
    ]
    ws.append(headers)

    def process_livraison(livraison):
        """Ajoute une ligne au tableau Excel à partir d'une livraison dict"""
        client = livraison.get("client", {})
        ws.append([
            client.get("nom", ""),
            client.get("telephone", ""),
            client.get("ville", ""),
            client.get("quartier", ""),
            livraison.get("montant_total", 0),
            livraison.get("montant_paye", 0),
            livraison.get("montant_restant", 0),
            livraison.get("quantite", 0),
            livraison.get("statut_paiement", ""),
            livraison.get("statut_livraison", ""),
            livraison.get("date_commande", ""),
            livraison.get("created_at", ""),
        ])

    def flatten_and_process(obj):
        """Gère récursivement les cas de liste imbriquée ou dict"""
        if isinstance(obj, list):
            for item in obj:
                flatten_and_process(item)
        elif isinstance(obj, dict):
            process_livraison(obj)
        else:
            print("⚠️ Élément inattendu ignoré :", repr(obj))

    flatten_and_process(livraisons)

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    return file_stream

#__________________________________________________________
#                                                          |
# Exportations des données du livreur sous format excel    |
#__________________________________________________________|







#__________________________________________________________
#                                                          |
# Exportations des données du staff sous format excel      |
#__________________________________________________________|

def export_magasins_to_excel(magasins):
    from openpyxl import Workbook
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = "Magasins"

    headers = [
        "Nom Magasin", "Adresse", "Téléphone", "Date de création",
        "Nombre d'utilisateurs", "Utilisateurs (noms)"
    ]
    ws.append(headers)

    for m in magasins:
        utilisateurs = m.get("utilisateurs", [])
        noms_utilisateurs = ", ".join([u.get("first_name", "") + " " + u.get("last_name", "") for u in utilisateurs])
        
        ws.append([
            m.get("nom", ""),
            m.get("adresse", ""),
            m.get("telephone", ""),
            m.get("created_at").strftime("%Y-%m-%d %H:%M") if isinstance(m.get("created_at"), datetime) else "",
            len(utilisateurs),
            noms_utilisateurs
        ])

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    return file_stream
    